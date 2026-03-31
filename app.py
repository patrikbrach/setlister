import io
import time

import requests
import streamlit as st
import openpyxl

from excel_reader import parse_excel
from setlist_client import fetch_setlist_sync
from pdf_generator import generate_setlist_pdf, safe_filename

st.set_page_config(page_title="Setlist Lookup", page_icon="🎵", layout="centered")
st.title("🎵 Setlist Lookup")

# ── API key ────────────────────────────────────────────────
api_key = st.text_input("setlist.fm API-nyckel", type="password", placeholder="Din API-nyckel")

# ── File upload ────────────────────────────────────────────
uploaded = st.file_uploader("Ladda upp Excel-fil (.xlsx)", type=["xlsx"])

run = st.button("Hämta setlister", disabled=not (api_key and uploaded))

if not run:
    st.stop()

# ── Parse Excel ────────────────────────────────────────────
try:
    rows = parse_excel(uploaded.read())
except ValueError as e:
    st.error(str(e))
    st.stop()

st.write(f"Hittade **{len(rows)}** rader. Hämtar setlister…")
progress = st.progress(0)
status = st.empty()

results = []

# ── Fetch each row ─────────────────────────────────────────
placeholders = [st.empty() for _ in rows]

for i, row in enumerate(rows):
    status.text(f"{i} / {len(rows)} klara")

    result = fetch_setlist_sync(
        api_key=api_key,
        row_id=row["id"],
        artist=row["artist"],
        date=row["date"],
        venue=row.get("venue"),
    )
    results.append(result)

    # Render result card
    with placeholders[i].container():
        if result["status"] == "found":
            songs = result.get("songs", [])
            encore_breaks = set(result.get("encore_breaks", []))
            label = f"✅ {result['artist']} – {result['date']} – {result.get('venue', '')}"
            with st.expander(label):
                lines = []
                for j, song in enumerate(songs):
                    if j in encore_breaks:
                        lines.append("— *Encore* —")
                    info = f" *({song['info']})*" if song.get("info") else ""
                    lines.append(f"{song['order']}. {song['name']}{info}")
                st.markdown("\n\n".join(lines) if lines else "*Tom setlist*")
                if result.get("setlist_url"):
                    st.markdown(f"[Öppna på setlist.fm]({result['setlist_url']})")
                pdf_bytes = generate_setlist_pdf(result)
                st.download_button(
                    label="⬇ Ladda ner som PDF",
                    data=pdf_bytes,
                    file_name=safe_filename(result["artist"], result["date"]),
                    mime="application/pdf",
                    key=f"pdf_{result['id']}",
                )

        elif result["status"] == "not_found":
            st.warning(f"❌ {result['artist']} – {result['date']} – Ingen setlist hittades")

        else:
            msg = result.get("message", "Okänt fel")
            st.error(f"⚠️ {result.get('artist', '')} – {result['date']} – {msg}")
            if "401" in msg:
                status.text("Ogiltig API-nyckel – avbryter.")
                progress.progress(1.0)
                st.stop()

    progress.progress((i + 1) / len(rows))

    if i < len(rows) - 1:
        time.sleep(0.5)

status.text(f"✅ {len(rows)} / {len(rows)} klara")

# ── Export ─────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Setlists"
ws.append(["Artist", "Datum", "Venue", "Antal låtar", "Setlist.fm URL"])
for r in results:
    ws.append([
        r.get("artist", ""),
        r.get("date", ""),
        r.get("venue", ""),
        len(r.get("songs", [])) if r["status"] == "found" else 0,
        r.get("setlist_url", ""),
    ])
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)

st.download_button(
    label="⬇ Exportera till Excel",
    data=buf,
    file_name="setlists.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
